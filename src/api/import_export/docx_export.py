from __future__ import annotations

from io import BytesIO

import html
import re
from pathlib import Path
from bs4 import BeautifulSoup, NavigableString
from fastapi import APIRouter
from fastapi.responses import FileResponse

from backend.environment import _HOSTED_LIMITS, runtime
from backend.data_access import allrows, one, project
from storage import context as storage_context
from storage import mutations as storage_mutations

router=APIRouter()

def _manuscript_block_groups(blocks:list[dict]):
    """Yield one prose/image block or one validated adjacent image row at a time."""
    ordered=list(blocks);i=0
    while i<len(ordered):
        block=ordered[i];row_id=str(block.get('image_row_id') or '').strip() if block.get('block_type')=='image' else ''
        if row_id:
            group=[block];j=i+1
            while j<len(ordered) and ordered[j].get('block_type')=='image' and str(ordered[j].get('image_row_id') or '').strip()==row_id:
                group.append(ordered[j]);j+=1
            if 2<=len(group)<=4:
                yield group;i=j;continue
        yield [block];i+=1

@router.get("/api/export/docx")
def export_docx():
    # Keep python-docx/lxml off Breadmaker's normal server startup path. Export is
    # an explicit action and may pay this dependency cost when invoked.
    from docx import Document
    with storage_mutations.tx() as con:
        p=project(con); tabs=allrows(con,"SELECT * FROM document_tabs WHERE project_id=? ORDER BY position,id",(p['id'],)); chapters=allrows(con,"SELECT * FROM chapters WHERE project_id=? ORDER BY tab_id,position,id",(p['id'],)); blocks=allrows(con,"SELECT * FROM manuscript_blocks WHERE project_id=? ORDER BY tab_id,chapter_id,position,id",(p['id'],))
    doc=Document(); doc.core_properties.title=p['title']; by_ch={c['id']:[] for c in chapters}
    for b in blocks: by_ch.setdefault(b['chapter_id'],[]).append(b)
    for ti,tab in enumerate(tabs):
        tabchs=[c for c in chapters if c['tab_id']==tab['id']]
        # Non-manuscript source tabs get a real Heading 1 so their information is retained. Ghost scene labels are never exported.
        if len(tabs)>1 and not tab['is_manuscript']: doc.add_heading(tab['title'],level=1)
        wrote_segment=False
        for ch in tabchs:
            chapter_blocks=by_ch.get(ch['id'],[])
            # The quiet persistence container is not a user-facing Chapter. Its
            # initial empty Body paragraph should not create a blank export page.
            if ch.get('is_implicit') and not any(b.get('block_type')=='image' or str(b.get('text_content') or '').strip() for b in chapter_blocks):
                continue
            if wrote_segment: doc.add_page_break()
            if not ch.get('is_implicit'):
                doc.add_heading(ch['title'],level=1)
            for group in _manuscript_block_groups(chapter_blocks):
                if len(group)>1 and all(b.get('block_type')=='image' and b.get('image_path') for b in group):
                    try:
                        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
                        from docx.enum.text import WD_ALIGN_PARAGRAPH
                        section=doc.sections[-1];usable_width=max(1,int(section.page_width)-int(section.left_margin)-int(section.right_margin))
                        row_width=max(5,min(100,int(group[0].get('image_width') or 100)));total_width=int(usable_width*row_width/100);cell_width=max(1,int(total_width/len(group)))
                        table=doc.add_table(rows=1,cols=len(group));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
                        for index,b in enumerate(group):
                            cell=table.cell(0,index);cell.width=cell_width;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
                            paragraph=cell.paragraphs[0];paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
                            asset_name=Path(str(b.get('image_path') or '')).name
                            if runtime.assets.exists(asset_name):
                                if _HOSTED_LIMITS is not None:
                                    paragraph.add_run('[Image]')
                                else:
                                    shape=paragraph.add_run().add_picture(BytesIO(runtime.assets.read_blob(asset_name)),width=max(1,int(cell_width*.90)))
                                    description=str(b.get('image_description') or '').strip()
                                    if description:
                                        try: shape._inline.docPr.set('descr',description)
                                        except Exception: pass
                            footer=str(b.get('image_footer') or '').strip()
                            if footer:
                                footer_paragraph=cell.add_paragraph(footer);footer_paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    except Exception: pass
                    continue
                b=group[0]
                if b['block_type']=='image' and b['image_path']:
                    asset_name=Path(b['image_path']).name
                    if runtime.assets.exists(asset_name):
                        try:
                            from docx.enum.text import WD_ALIGN_PARAGRAPH
                            section=doc.sections[-1]
                            usable_width=max(1,int(section.page_width)-int(section.left_margin)-int(section.right_margin))
                            width_percent=max(5,min(100,int(b.get('image_width') or 100)))
                            image_paragraph=doc.add_paragraph();image_paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
                            if _HOSTED_LIMITS is not None:
                                image_paragraph.add_run('[Image]')
                            else:
                                shape=image_paragraph.add_run().add_picture(BytesIO(runtime.assets.read_blob(asset_name)),width=int(usable_width*width_percent/100))
                                description=str(b.get('image_description') or '').strip()
                                if description:
                                    try: shape._inline.docPr.set('descr',description)
                                    except Exception: pass
                            footer=str(b.get('image_footer') or '').strip()
                            if footer:
                                footer_paragraph=doc.add_paragraph(footer);footer_paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
                                footer_paragraph.paragraph_format.left_indent=int(usable_width*.07);footer_paragraph.paragraph_format.right_indent=int(usable_width*.07)
                        except Exception: pass
                    continue
                if b['semantic_style'] in {'heading2','heading3','heading4'}: doc.add_heading(b['text_content'],level=int(b['semantic_style'][-1])); continue
                p_=doc.add_paragraph(); soup=BeautifulSoup(b['content_html'] or html.escape(b['text_content']),'html.parser')
                def walk(node,bold=False,italic=False,underline=False):
                    if isinstance(node,NavigableString):
                        text=str(node); parts=text.split('\t')
                        for i,part in enumerate(parts):
                            if part:
                                r=p_.add_run(part); r.bold=bold; r.italic=italic; r.underline=underline
                            if i<len(parts)-1: p_.add_run('\t')
                    else:
                        nb=bold or node.name in {'b','strong'}; ni=italic or node.name in {'i','em'}; nu=underline or node.name=='u'
                        for child in node.children: walk(child,nb,ni,nu)
                for child in soup.contents: walk(child)
            wrote_segment=True
        if ti!=len(tabs)-1: doc.add_page_break()
    safe=re.sub(r'[^A-Za-z0-9._-]+','_',p['title']).strip('_') or 'book'; path=(storage_context.current_data_dir()/'exports')/f'{safe}.docx'; path.parent.mkdir(parents=True,exist_ok=True); doc.save(path); return FileResponse(path,filename=path.name,media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
