from __future__ import annotations

"""Transient document import/export helpers for the M5 hosted runtime.

All parsing/rendering happens in request memory. Canonical content is committed only
through HostedDomainStore/AssetStore; this module never creates a durable server file.
"""

import html
import io
import re
import zipfile
from typing import Any, Mapping

from docx import Document
from docx.shared import Inches
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.lib.pagesizes import LETTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image

_MAX_DOCUMENT_BYTES=64*1024*1024
_MAX_DOCX_ENTRIES=10000
_MAX_DOCX_UNCOMPRESSED=256*1024*1024


def read_upload_bytes_sync(data:bytes) -> bytes:
    payload=bytes(data)
    if len(payload)>_MAX_DOCUMENT_BYTES:raise ValueError('HOSTED_DOCUMENT_TOO_LARGE')
    return payload


def validate_docx_bytes(data:bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data),'r') as archive:
            infos=archive.infolist()
            if len(infos)>_MAX_DOCX_ENTRIES:raise ValueError('HOSTED_DOCX_ARCHIVE_TOO_LARGE')
            total=0
            for info in infos:
                name=info.filename.replace('\\','/')
                if name.startswith('/') or '..' in name.split('/'):raise ValueError('HOSTED_DOCX_PATH_INVALID')
                total+=int(info.file_size)
                if total>_MAX_DOCX_UNCOMPRESSED:raise ValueError('HOSTED_DOCX_ARCHIVE_TOO_LARGE')
            if '[Content_Types].xml' not in {i.filename for i in infos}:raise ValueError('HOSTED_DOCX_INVALID')
    except zipfile.BadZipFile as exc:raise ValueError('HOSTED_DOCX_INVALID') from exc


def parse_document_bytes(data:bytes, filename:str, chapter_heading:str='heading1') -> tuple[list[dict[str,Any]],list[str]]:
    payload=read_upload_bytes_sync(data);lower=str(filename or 'document.docx').lower();warnings=[]
    if lower.endswith('.txt'):
        text=payload.decode('utf-8',errors='replace')
        blocks=[{'kind':'paragraph','text':line,'html':html.escape(line),'style':'body'} for line in text.splitlines()]
        return [{'title':'Manuscript','chapters':[{'title':'Chapter 1','blocks':blocks or [{'kind':'paragraph','text':'','html':'','style':'body'}]}]}],warnings
    if lower.endswith('.doc'):
        raise ValueError('HOSTED_LEGACY_DOC_DISABLED')
    if not lower.endswith('.docx'):raise ValueError('HOSTED_DOCUMENT_EXTENSION_UNSUPPORTED')
    validate_docx_bytes(payload)
    try:document=Document(io.BytesIO(payload))
    except Exception as exc:raise ValueError('HOSTED_DOCX_INVALID') from exc
    level_match=re.fullmatch(r'heading([1-6])',str(chapter_heading or '').strip().lower())
    if not level_match:raise ValueError('HOSTED_CHAPTER_HEADING_INVALID')
    chapter_level=int(level_match.group(1))
    chapters=[];current={'title':'Chapter 1','blocks':[]}
    def push():
        nonlocal current
        if current['blocks'] or current['title']!='Chapter 1':chapters.append(current)
        current={'title':f'Chapter {len(chapters)+1}','blocks':[]}
    for paragraph in document.paragraphs:
        style=str(getattr(paragraph.style,'name','') or '').strip().lower()
        match=re.fullmatch(r'heading\s*([1-6])',style)
        text=str(paragraph.text or '')
        if match and int(match.group(1))==chapter_level:
            if current['blocks'] or chapters:push()
            current['title']=text.strip()[:240] or f'Chapter {len(chapters)+1}'
            continue
        if any('drawing' in child.tag.lower() or 'pict' in child.tag.lower() for run in paragraph.runs for child in run._r):
            warnings.append('Embedded DOCX images were omitted during hosted text import; add them again from the manuscript image tool.')
        current['blocks'].append({'kind':'paragraph','text':text,'html':html.escape(text),'style':'body'})
    if current['blocks'] or not chapters:chapters.append(current)
    for chapter in chapters:
        if not chapter['blocks']:chapter['blocks']=[{'kind':'paragraph','text':'','html':'','style':'body'}]
    return [{'title':'Manuscript','chapters':chapters}],list(dict.fromkeys(warnings))


def ordered_book(domain,project_id:str,user_id:str) -> tuple[dict[str,Any],list[dict[str,Any]],dict[str,list[dict[str,Any]]]]:
    tabs=sorted(domain.list('document_tab',{'project_id':project_id},user_id=user_id),key=lambda row:(int(row.get('position') or 0),str(row.get('id') or '')))
    tab=next((row for row in tabs if row.get('is_manuscript') or str(row.get('kind') or '')=='manuscript'),tabs[0] if tabs else {'id':''})
    chapters=sorted([row for row in domain.list('chapter',{'project_id':project_id},user_id=user_id) if str(row.get('tab_id') or '')==str(tab.get('id') or '')],key=lambda row:(int(row.get('position') or 0),str(row.get('id') or '')))
    blocks=domain.list('manuscript_block',{'project_id':project_id},user_id=user_id)
    grouped={str(ch.get('id')):sorted([b for b in blocks if str(b.get('chapter_id') or '')==str(ch.get('id'))],key=lambda row:(float(row.get('position') or 0),str(row.get('id') or ''))) for ch in chapters}
    return tab,chapters,grouped


def export_docx_bytes(domain,assets,project:Mapping[str,Any],project_id:str,user_id:str) -> bytes:
    _tab,chapters,grouped=ordered_book(domain,project_id,user_id);doc=Document();title=str(project.get('title') or project.get('name') or 'Untitled')
    doc.add_heading(title,0)
    for chapter in chapters:
        doc.add_heading(str(chapter.get('title') or 'Chapter'),level=1)
        for block in grouped.get(str(chapter.get('id')),[]):
            if str(block.get('block_type') or 'paragraph')=='image' and block.get('image_path'):
                try:
                    payload=assets.read_blob(str(block['image_path']));width=max(0.5,min(6.5,6.5*int(block.get('image_width') or 100)/100));doc.add_picture(io.BytesIO(payload),width=Inches(width))
                    footer=str(block.get('image_footer') or '').strip()
                    if footer:doc.add_paragraph(footer,style='Caption')
                except Exception:
                    doc.add_paragraph('[Image unavailable]')
            else:
                p=doc.add_paragraph(str(block.get('text_content') or ''))
                style=str(block.get('semantic_style') or '').lower()
                if style.startswith('heading'):
                    try:p.style='Heading '+str(max(1,min(9,int(re.sub(r'\D','',style) or '2'))))
                    except Exception:pass
    out=io.BytesIO();doc.save(out);return out.getvalue()


def export_pdf_bytes(domain,assets,project:Mapping[str,Any],project_id:str,user_id:str) -> bytes:
    _tab,chapters,grouped=ordered_book(domain,project_id,user_id);out=io.BytesIO();styles=getSampleStyleSheet();story=[]
    title_style=ParagraphStyle('BreadmakerTitle',parent=styles['Title'],alignment=TA_CENTER,spaceAfter=18)
    story.append(Paragraph(html.escape(str(project.get('title') or project.get('name') or 'Untitled')),title_style))
    for chapter in chapters:
        story.append(Paragraph(html.escape(str(chapter.get('title') or 'Chapter')),styles['Heading1']))
        for block in grouped.get(str(chapter.get('id')),[]):
            if str(block.get('block_type') or 'paragraph')=='image' and block.get('image_path'):
                try:
                    payload=assets.read_blob(str(block['image_path']));reader=ImageReader(io.BytesIO(payload));w,h=reader.getSize();max_w=430;scale=min(1,max_w/max(1,w));story.append(Image(io.BytesIO(payload),width=w*scale,height=h*scale));story.append(Spacer(1,6))
                    footer=str(block.get('image_footer') or '').strip()
                    if footer:story.append(Paragraph(html.escape(footer),styles['Italic']))
                except Exception:story.append(Paragraph('[Image unavailable]',styles['Italic']))
            else:
                text=html.escape(str(block.get('text_content') or '')).replace('\n','<br/>');story.append(Paragraph(text or '&nbsp;',styles['BodyText']));story.append(Spacer(1,6))
    SimpleDocTemplate(out,pagesize=LETTER,rightMargin=54,leftMargin=54,topMargin=54,bottomMargin=54,title=str(project.get('title') or 'Breadmaker')).build(story)
    return out.getvalue()


def manuscript_is_default_blank(domain,project_id:str,user_id:str) -> bool:
    tab,chapters,grouped=ordered_book(domain,project_id,user_id)
    if not tab or len(chapters)!=1:return False
    blocks=grouped.get(str(chapters[0].get('id')),[])
    return len(blocks)==1 and str(blocks[0].get('block_type') or 'paragraph')=='paragraph' and not str(blocks[0].get('text_content') or '').strip() and not str(blocks[0].get('image_path') or '').strip()


def install_parsed_tabs(domain,project_id:str,user_id:str,tabs:list[dict[str,Any]],*,replace_default:bool=False,after_document_id:str='') -> dict[str,Any]:
    if not tabs:return {'document_ids':[],'chapters':0,'images':0}
    existing_tabs=sorted(domain.list('document_tab',{'project_id':project_id},user_id=user_id),key=lambda row:(int(row.get('position') or 0),str(row.get('id') or '')))
    document_ids=[];chapter_count=0
    for tab_index,parsed_tab in enumerate(tabs):
        if replace_default and tab_index==0:
            tab=next((row for row in existing_tabs if row.get('is_manuscript') or str(row.get('kind') or '')=='manuscript'),existing_tabs[0])
            chapters=sorted([row for row in domain.list('chapter',{'project_id':project_id},user_id=user_id) if str(row.get('tab_id') or '')==str(tab['id'])],key=lambda row:int(row.get('position') or 0))
            base_chapter=chapters[0]
            old_blocks=[row for row in domain.list('manuscript_block',{'project_id':project_id},user_id=user_id) if str(row.get('chapter_id') or '')==str(base_chapter['id'])]
            parsed_chapters=parsed_tab.get('chapters') or [{'title':'Chapter 1','blocks':[]}]
            first=parsed_chapters[0];domain.update('chapter',{'project_id':project_id,'chapter_id':str(base_chapter['id'])},{'title':str(first.get('title') or 'Chapter 1'),'position':0},user_id=user_id)
            first_blocks=first.get('blocks') or [{'text':'','html':'','style':'body'}]
            domain.update('manuscript_block',{'project_id':project_id,'block_id':str(old_blocks[0]['id'])},{'text_content':str(first_blocks[0].get('text') or ''),'content_html':str(first_blocks[0].get('html') or ''),'semantic_style':str(first_blocks[0].get('style') or 'body'),'position':0},user_id=user_id)
            for block in old_blocks[1:]:domain.delete('manuscript_block',{'project_id':project_id,'block_id':str(block['id'])},user_id=user_id)
            for pos,block in enumerate(first_blocks[1:],1):domain.create('manuscript_block',{'project_id':project_id},{'tab_id':tab['id'],'chapter_id':base_chapter['id'],'position':pos,'block_type':'paragraph','semantic_style':str(block.get('style') or 'body'),'text_content':str(block.get('text') or ''),'content_html':str(block.get('html') or '')},user_id=user_id)
            chapter_count+=1;document_ids.append(str(tab['id']))
            for cpos,chapter in enumerate(parsed_chapters[1:],1):
                ch=domain.create('chapter',{'project_id':project_id},{'tab_id':tab['id'],'title':str(chapter.get('title') or f'Chapter {cpos+1}'),'position':cpos},user_id=user_id);chapter_count+=1
                for bpos,block in enumerate(chapter.get('blocks') or [{'text':'','html':'','style':'body'}]):domain.create('manuscript_block',{'project_id':project_id},{'tab_id':tab['id'],'chapter_id':ch['id'],'position':bpos,'block_type':'paragraph','semantic_style':str(block.get('style') or 'body'),'text_content':str(block.get('text') or ''),'content_html':str(block.get('html') or '')},user_id=user_id)
            continue
        tab=domain.create('document_tab',{'project_id':project_id},{'title':str(parsed_tab.get('title') or f'Document {len(existing_tabs)+tab_index+1}')[:240],'kind':'document','position':len(existing_tabs)+tab_index},user_id=user_id);document_ids.append(str(tab['id']))
        for cpos,chapter in enumerate(parsed_tab.get('chapters') or [{'title':'Chapter 1','blocks':[]}]):
            ch=domain.create('chapter',{'project_id':project_id},{'tab_id':tab['id'],'title':str(chapter.get('title') or f'Chapter {cpos+1}')[:240],'position':cpos},user_id=user_id);chapter_count+=1
            for bpos,block in enumerate(chapter.get('blocks') or [{'text':'','html':'','style':'body'}]):domain.create('manuscript_block',{'project_id':project_id},{'tab_id':tab['id'],'chapter_id':ch['id'],'position':bpos,'block_type':'paragraph','semantic_style':str(block.get('style') or 'body'),'text_content':str(block.get('text') or ''),'content_html':str(block.get('html') or '')},user_id=user_id)
    return {'document_ids':document_ids,'chapters':chapter_count,'images':0}
