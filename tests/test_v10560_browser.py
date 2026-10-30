"""Cold book acceleration and speculative metadata never bypass editing gates."""
import io
import re
from docx import Document
from playwright.sync_api import expect
from tests.test_v10548_browser import hosted_browser


def project(base,page):
    doc=Document()
    for n in range(6):
        doc.add_heading(f'Synthetic60 chapter {n}',level=1)
        for i in range(100):doc.add_paragraph(f'Synthetic60 chapter {n} paragraph {i}')
    raw=io.BytesIO();doc.save(raw);headers={'Origin':base}
    job=page.request.post(base+'/api/import/new-project?resumable=true&project_name=Synthetic60',headers=headers,
        multipart={'file':{'name':'synthetic60.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','buffer':raw.getvalue()}}).json()
    while not job['ready']:
        job=page.request.post(base+f'/api/import/jobs/{job["job_id"]}/step?batch={job["next_batch"]}',headers=headers).json()
    pid=page.request.post(base+f'/api/import/jobs/{job["job_id"]}/finish',headers=headers).json()['project']['id']
    page.goto(base+'/p/'+pid,wait_until='domcontentloaded');page.wait_for_function('()=>document.documentElement.dataset.breadmakerBootReady==="true"')
    return pid


def test_fast_full_text_shows_existing_readonly_text_and_requires_final_authority(tmp_path):
    with hosted_browser(tmp_path) as (base,page):
        project(base,page);first=[];final=[];bodies=[]
        page.get_by_text('Synthetic60 chapter 5',exact=True).first.click()
        page.wait_for_function('()=>state.chapter?.title==="Synthetic60 chapter 5"')
        page.on('request',lambda r:bodies.append(r.url) if '/api/chapters/' in r.url and r.method=='GET' else None)
        def hold(route):
            if 'bundled=true' in route.request.url:
                first.append(route.request.url)
                expect(page.locator('.full-text-preview-body')).to_contain_text('Synthetic60 chapter 5 paragraph 0')
                assert page.locator('.full-text-preview .book-chapter-heading').count()==1
                assert page.locator('.full-text-preview').get_attribute('contenteditable')=='false'
            elif route.request.headers.get('if-none-match'):
                final.append(route.request.url)
                assert page.evaluate('()=>state.bookCache.size===0&&state.viewMode==="chapter"')
                assert page.locator('.full-text-preview').get_attribute('contenteditable')=='false'
            route.continue_()
        page.route('**/api/book?**',hold)
        page.locator('#bookViewBtn').click()
        page.wait_for_function('()=>state.viewMode==="book"&&!fullTextLoading')
        assert len(first)==len(final)==1
        assert page.locator('#editor .manuscript-block').count()==600
        assert not bodies


def test_unavailable_book_accelerator_falls_back_to_complete_verified_chapters(tmp_path):
    with hosted_browser(tmp_path) as (base,page):
        project(base,page);bodies=[]
        page.on('request',lambda r:bodies.append(r.url) if '/api/chapters/' in r.url and r.method=='GET' else None)
        def unavailable(route):
            if 'bundled=true' not in route.request.url:route.continue_();return
            response=route.fetch();data=response.json();data.update(blocks=[],bundled=False)
            route.fulfill(response=response,json=data)
        page.route('**/api/book?**',unavailable)
        page.locator('#bookViewBtn').click();page.wait_for_function('()=>state.viewMode==="book"&&!fullTextLoading')
        assert page.locator('#editor .manuscript-block').count()==600
        assert len(bodies)==5 and page.evaluate('()=>state.bookData.blocks.length')==600


def test_changed_final_generation_cannot_publish_fast_book_or_unlock_stale_preview(tmp_path):
    with hosted_browser(tmp_path) as (base,page):
        project(base,page)
        def changed(route):
            if route.request.headers.get('if-none-match'):route.fulfill(status=200,json={'project':{'generation':999999}})
            else:route.continue_()
        page.route('**/api/book?**',changed)
        page.locator('#bookViewBtn').click();page.wait_for_function('()=>!fullTextLoading')
        assert page.evaluate('()=>state.viewMode==="chapter"&&state.bookCache.size===0&&!state.dirty')
        assert page.locator('#editor .manuscript-block').count()==100
        expect(page.locator('.full-text-preview')).to_have_count(0)


def test_routed_metadata_overlaps_workspace_but_never_paints_before_workspace_and_lease(tmp_path):
    with hosted_browser(tmp_path) as (base,page):
        pid=project(base,page);workspace=[];bootstrap=[];leases=[]
        page.route('**/api/workspace',lambda r:workspace.append(r))
        page.on('request',lambda r:bootstrap.append(r) if '/api/bootstrap?' in r.url else leases.append(r) if r.url.endswith('/api/lease') else None)
        with page.expect_response(lambda r:'/api/bootstrap?' in r.url):
            with page.expect_request('**/api/workspace'):page.reload(wait_until='domcontentloaded')
        assert workspace and bootstrap
        assert bootstrap[0].headers['x-breadmaker-project']==pid
        assert not leases
        expect(page.locator('#startupScreen')).to_be_visible();expect(page.locator('#app')).to_be_hidden()
        workspace[0].continue_();page.wait_for_function('()=>document.documentElement.dataset.breadmakerBootReady==="true"')
        expect(page.locator('#app')).to_be_visible();assert leases


def test_missing_workspace_membership_redirects_without_installing_speculative_metadata(tmp_path):
    with hosted_browser(tmp_path) as (base,page):
        original=project(base,page)
        first=[True];creates=[]
        def missing_membership(route):
            if first[0]:
                first[0]=False;route.fulfill(status=200,json={'projects':[],'runtime':'hosted_dev','active_project_id':''})
            else:route.continue_()
        page.route('**/api/workspace',missing_membership)
        page.on('request',lambda r:creates.append(r) if r.method=='POST' and r.url==base+'/api/projects' else None)
        page.reload(wait_until='domcontentloaded')
        page.wait_for_url(re.compile(r'/p/(?!'+original+r'$)[0-9a-f]{32}$'))
        page.wait_for_function('()=>document.documentElement.dataset.breadmakerBootReady==="true"')
        expect(page.locator('.manuscript-block')).to_have_count(1)
        expect(page.locator('.manuscript-block').first).to_have_text('')
        expect(page.locator('#saveState')).to_have_text('Saved')
        assert page.evaluate('()=>state.workspace.active_project_id')!=original
        assert len(creates)==1


def test_denied_fast_book_cannot_leave_retained_text_editable_or_cache_reply(tmp_path):
    with hosted_browser(tmp_path) as (base,page):
        project(base,page)
        page.route('**/api/book?**',lambda r:r.fulfill(status=403,json={'detail':'Synthetic revoked membership'}))
        page.locator('#bookViewBtn').click();page.wait_for_function('()=>!fullTextLoading')
        assert page.evaluate('()=>state.viewMode==="chapter"&&state.bookCache.size===0&&state.chapterCache.size===0&&!state.bookDOMActive&&!state.bookDOMCache')
        assert page.locator('#manuscriptPrimaryPane').evaluate('(el)=>el.inert')
        expect(page.locator('.full-text-preview')).to_have_count(0)


def test_logout_during_fast_book_final_check_cannot_repopulate_memory(tmp_path):
    with hosted_browser(tmp_path) as (base,page):
        project(base,page);held=[]
        def hold(route):
            if route.request.headers.get('if-none-match'):held.append(route)
            else:route.continue_()
        page.route('**/api/book?**',hold)
        with page.expect_request(lambda r:'/api/book?' in r.url and bool(r.headers.get('if-none-match'))):
            page.locator('#bookViewBtn').click()
        page.route('**/api/auth/logout',lambda r:r.fulfill(status=200,json={'ok':True}))
        page.evaluate('async()=>await api("/api/auth/logout",{method:"POST"})')
        assert held;held[0].continue_();page.wait_for_function('()=>!fullTextLoading')
        assert page.evaluate('()=>state.viewMode==="chapter"&&state.bookCache.size===0&&state.chapterCache.size===0')
