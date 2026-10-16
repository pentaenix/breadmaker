from __future__ import annotations
from source_contracts import frontend_source_from_app_path,frontend_styles_source_from_styles_path

from api.import_export import docx_parser as import_api

import json
from pathlib import Path
from tempfile import NamedTemporaryFile

from docx import Document

from conftest import assert_ok


def test_v41_config_metadata_and_connector_ports_persist(client, fresh_project):
    tab=assert_ok(client.post('/api/documents',json={'title':'Research','description':'Worldbuilding notes'}))
    assert_ok(client.patch(f"/api/documents/{tab['id']}",json={'description':'People, places, and rules'}))

    canvas=assert_ok(client.post('/api/canvases',json={'name':'Arc Map','description':'Cinna development','canvas_type':'blank'}))
    assert canvas['description']=='Cinna development'
    assert_ok(client.patch(f"/api/canvases/{canvas['id']}",json={'description':'Cinna and Bohor'}))

    story_map=assert_ok(client.post('/api/maps',json={'name':'Main Story','description':'Primary chronology'}))
    assert story_map['description']=='Primary chronology'
    assert_ok(client.patch(f"/api/maps/{story_map['id']}",json={'description':'Primary plot chronology'}))

    first=assert_ok(client.post(f"/api/canvases/{canvas['id']}/nodes",json={'title':'A','x':80,'y':80,'w':180,'h':100}))
    second=assert_ok(client.post(f"/api/canvases/{canvas['id']}/nodes",json={'title':'B','x':420,'y':260,'w':180,'h':100}))
    edge=assert_ok(client.post(f"/api/canvases/{canvas['id']}/edges",json={
        'from_node_id':first['id'],'to_node_id':second['id'],'from_port':'bottom_right','to_port':'top_left','label':'turn'
    }))
    assert edge['from_port']=='bottom_right' and edge['to_port']=='top_left'
    assert_ok(client.patch(f"/api/edges/{edge['id']}",json={'from_port':'top_right','to_port':'bottom_left','label':'revised'}))

    boot=assert_ok(client.get('/api/bootstrap'))
    saved_tab=next(d for d in boot['tabs'] if d['id']==tab['id'])
    saved_canvas=next(c for c in boot['canvases'] if c['id']==canvas['id'])
    saved_map=next(m for m in boot['maps'] if m['id']==story_map['id'])
    saved_edge=assert_ok(client.get(f"/api/canvases/{canvas['id']}"))['edges'][0]
    assert saved_tab['description']=='People, places, and rules'
    assert saved_canvas['description']=='Cinna and Bohor'
    assert saved_map['description']=='Primary plot chronology'
    assert (saved_edge['from_port'],saved_edge['to_port'],saved_edge['label'])==('top_right','bottom_left','revised')


def test_v41_semantic_heading_levels_are_bounded_to_three_sublevels(client, fresh_project):
    chapter=fresh_project['bootstrap']['chapters'][0]
    blocks=[
        {'block_type':'paragraph','semantic_style':'heading3','text_content':'Scene heading','content_html':'Scene heading'},
        {'block_type':'paragraph','semantic_style':'heading6','text_content':'Small heading','content_html':'Small heading'},
        {'block_type':'paragraph','semantic_style':'body','text_content':'Prose','content_html':'Prose'},
    ]
    assert_ok(client.put(f"/api/chapters/{chapter['id']}/blocks",json={'blocks':blocks}))
    saved=assert_ok(client.get(f"/api/chapters/{chapter['id']}"))['blocks']
    assert [b['semantic_style'] for b in saved]==['heading3','body','body']


def test_v41_docx_chapter_marker_is_configurable(app_module_fixture):
    doc=Document()
    doc.add_paragraph('Preface heading',style='Heading 1')
    doc.add_paragraph('Chapter A',style='Heading 2')
    doc.add_paragraph('Internal section',style='Heading 3')
    doc.add_paragraph('A prose paragraph')
    doc.add_paragraph('Chapter B',style='Heading 2')
    doc.add_paragraph('More prose')
    with NamedTemporaryFile(suffix='.docx',delete=False) as tmp:
        path=Path(tmp.name)
    try:
        doc.save(path)
        tabs=import_api.parse_docx(path,'heading2')
        chapters=tabs[0]['chapters']
        assert [c['title'] for c in chapters]==['Chapter 1','Chapter A','Chapter B']
        preface=chapters[0]['blocks'][0]
        internal=next(b for b in chapters[1]['blocks'] if b.get('text')=='Internal section')
        assert preface['style']=='body'
        assert internal['style']=='heading2'
    finally:
        path.unlink(missing_ok=True)


def test_v41_static_interaction_contracts():
    root=Path(__file__).resolve().parents[1]
    html=(root/'static/index.html').read_text(encoding='utf-8')
    js=frontend_source_from_app_path(root/'static/app.js')
    css=frontend_styles_source_from_styles_path((root/'static/styles.css'))
    locale=json.loads((root/'locales/en.json').read_text(encoding='utf-8'))
    strings=locale['strings']
    assert 'Tabs' in strings.values() and 'Full Text' in strings.values()
    assert '{{i18n:' in html
    assert 'id="findScope"' not in html
    assert 'id="focusExitBtn"' not in html
    assert 'syncFocusModeControl' in js and "$('#viewFocusBtn')?.addEventListener('click'" in js
    assert 'id="blockStyleBtn"' not in html and 'data-selection-action="style"' in html and 'heading3' in js and 'Heading 6' not in html
    assert 'MAP_ROUTE_WIDTH=10' in js and 'MAP_STOP_R=13' in js and 'MAP_STOP_INNER_R=8' in js
    assert 'renderMapSvg' in js and '/api/map-populate' in js
    assert "class','edge-hit'" in js
    assert 'bottom_right' in js and 'top_left' in js
    assert 'Move to Open' in js
    assert 'showNestedModal' in js
    assert '.brand-mark{width:27px;height:27px;border-radius:50%' in css
    assert '.paper{width:min(900px,100%);min-height:0' in css


def test_v41_revision_restore_preserves_new_configuration_metadata(client, fresh_project):
    tab=assert_ok(client.post('/api/documents',json={'title':'Notes','description':'Original tab description'}))
    canvas=assert_ok(client.post('/api/canvases',json={'name':'Plan','description':'Original canvas description','canvas_type':'blank'}))
    story_map=assert_ok(client.post('/api/maps',json={'name':'Arc','description':'Original map description'}))
    line=assert_ok(client.post(f"/api/maps/{story_map['id']}/lines",json={'name':'Character Arc'}))
    stop=assert_ok(client.post(f"/api/map-lines/{line['id']}/items",json={'title':'Beginning','planned_start':.1,'schematic_start':.1}))
    revision=assert_ok(client.post('/api/autocommit?label='+__import__('urllib.parse').parse.quote(str('V4.1 metadata'))))['id']

    assert_ok(client.patch(f"/api/documents/{tab['id']}",json={'description':'Changed'}))
    assert_ok(client.patch(f"/api/canvases/{canvas['id']}",json={'description':'Changed'}))
    assert_ok(client.patch(f"/api/maps/{story_map['id']}",json={'description':'Changed'}))
    assert_ok(client.patch(f"/api/map-items/{stop['id']}",json={'title':'Changed stop'}))
    assert_ok(client.post(f'/api/revisions/{revision}/restore'))

    boot=assert_ok(client.get('/api/bootstrap'))
    assert next(x for x in boot['tabs'] if x['id']==tab['id'])['description']=='Original tab description'
    assert next(x for x in boot['canvases'] if x['id']==canvas['id'])['description']=='Original canvas description'
    saved_map=next(x for x in boot['maps'] if x['id']==story_map['id'])
    assert saved_map['description']=='Original map description'
    restored=assert_ok(client.get(f"/api/maps/{story_map['id']}"))
    assert restored['lines'][0]['name']=='Character Arc'
    assert restored['lines'][0]['items'][0]['title']=='Beginning'
