from __future__ import annotations
from source_contracts import frontend_source_from_app_path,frontend_styles_source_from_styles_path

from api.import_export import docx_parser as import_api

import io
from pathlib import Path
from docx import Document
from conftest import assert_ok


def test_v410_static_outline_contract():
    root=Path(__file__).resolve().parents[1]
    html=(root/'static/index.html').read_text(encoding='utf-8')
    js=frontend_source_from_app_path(root/'static/app.js')
    css=frontend_styles_source_from_styles_path((root/'static/styles.css'))
    assert tuple(map(int,(root/'VERSION').read_text(encoding='utf-8').strip().split('.'))) >= (5, 0, 2)
    assert 'styles.css?v=5.0.2' in html and 'app.js?v=5.0.2' in html
    assert 'id="blockStyleBtn"' not in html and 'id="blockStylePopover"' not in html
    assert 'id="blockStyleSelect"' not in html
    assert 'data-selection-action="style"' in html
    assert "{label:'H1  Chapter',action:()=>applyWritingStyle('chapter')}" in js
    assert "{label:'H2  Section',action:()=>applyWritingStyle('heading2')}" in js
    assert "{label:'H3  Subsection',action:()=>applyWritingStyle('heading3')}" in js
    assert "{label:'H4  Minor section',action:()=>applyWritingStyle('heading4')}" in js
    assert "applyWritingStyle('heading1')" not in js
    assert 'chapter_outlines' in js
    assert '.outline-nav' in css and '.chapter-outline-toggle' in css
    assert '.outline-toggle{' not in css and '.book-chapter-toggle{' not in css


def test_v410_api_accepts_only_three_sublevels_below_chapter(client,fresh_project):
    chapter=fresh_project['bootstrap']['chapters'][0]
    assert_ok(client.put(f"/api/chapters/{chapter['id']}/blocks",json={'blocks':[
        {'semantic_style':'heading1','text_content':'Legacy H1 becomes body','content_html':'Legacy H1 becomes body'},
        {'semantic_style':'heading2','text_content':'A','content_html':'A'},
        {'semantic_style':'heading3','text_content':'B','content_html':'B'},
        {'semantic_style':'heading4','text_content':'C','content_html':'C'},
        {'semantic_style':'heading5','text_content':'Too deep','content_html':'Too deep'},
    ]}))
    saved=assert_ok(client.get(f"/api/chapters/{chapter['id']}"))['blocks']
    assert [b['semantic_style'] for b in saved]==['body','heading2','heading3','heading4','body']


def test_v410_docx_export_uses_chapter_h1_and_sublevels_h2_h3_h4(client,fresh_project):
    chapter=fresh_project['bootstrap']['chapters'][0]
    assert_ok(client.put(f"/api/chapters/{chapter['id']}/blocks",json={'blocks':[
        {'semantic_style':'heading2','text_content':'A','content_html':'A'},
        {'semantic_style':'heading3','text_content':'B','content_html':'B'},
        {'semantic_style':'heading4','text_content':'C','content_html':'C'},
        {'semantic_style':'body','text_content':'Prose','content_html':'Prose'},
    ]}))
    response=client.get('/api/export/docx');assert response.status_code==200
    doc=Document(io.BytesIO(response.content))
    rows=[(p.text,p.style.name) for p in doc.paragraphs if p.text]
    assert rows[:5]==[(chapter['title'],'Heading 1'),('A','Heading 2'),('B','Heading 3'),('C','Heading 4'),('Prose','Normal')]


def test_v410_docx_import_normalizes_headings_below_chapter_level(app_module_fixture,tmp_path):
    doc=Document();doc.add_paragraph('Chapter A',style='Heading 2');doc.add_paragraph('A',style='Heading 3');doc.add_paragraph('B',style='Heading 4');doc.add_paragraph('C',style='Heading 5');doc.add_paragraph('Too deep',style='Heading 6');doc.add_paragraph('Chapter B',style='Heading 2')
    path=tmp_path/'outline.docx';doc.save(path)
    tabs=import_api.parse_docx(path,'heading2');blocks=tabs[0]['chapters'][0]['blocks']
    assert [(b['text'],b['style']) for b in blocks]==[('A','heading2'),('B','heading3'),('C','heading4'),('Too deep','heading4')]
