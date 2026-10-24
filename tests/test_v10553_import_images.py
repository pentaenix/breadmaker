"""New-project DOCX images belong to the destination, not the source project."""
import io
import pytest

from docx import Document
from PIL import Image

from conftest import assert_ok


def document_with_images():
    document=Document();images=[]
    document.add_heading('Synthetic illustrated chapter',level=1)
    for index in range(5):
        out=io.BytesIO();Image.new('RGB',(12+index,12),(index*40,80,120)).save(out,format='PNG')
        blob=out.getvalue();images.append(blob)
        document.add_paragraph(f'Synthetic paragraph {index}')
        document.add_picture(io.BytesIO(blob))
    out=io.BytesIO();document.save(out)
    return out.getvalue(),images


def test_new_project_docx_images_are_durable_in_destination_only(client,fresh_project):
    source=assert_ok(client.get('/api/workspace'))['active_project_id']
    payload,images=document_with_images()
    result=assert_ok(client.post('/api/import/new-project',files={'file':('Synthetic illustrated.docx',payload,
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document')}))
    destination=result['project']['id']
    assert destination!=source and result['images']==5
    boot=assert_ok(client.get('/api/bootstrap'))
    book=assert_ok(client.get('/api/book',params={'tab_id':boot['tabs'][0]['id']}))
    image_blocks=[block for block in book['blocks'] if block['block_type']=='image']
    assert len(image_blocks)==5
    for block,blob in zip(image_blocks,images):
        response=client.get(block['image_path'])
        assert response.status_code==200 and response.content==blob
    integrity=assert_ok(client.get('/api/integrity/deep'))
    assert integrity['ok'],integrity
    assert_ok(client.post(f'/api/projects/{source}/activate'))
    for block in image_blocks:assert client.get(block['image_path']).status_code==404
    assert_ok(client.post(f'/api/projects/{destination}/activate'))
    for block,blob in zip(image_blocks,images):assert client.get(block['image_path']).content==blob


def test_hosted_docx_publishes_verified_assets_before_project_visibility(monkeypatch):
    from hosted.document_import import DocumentImport
    from hosted.firestore.contracts import ProjectUnavailable
    from hosted.runtime_services import _DEV_USER_ID as owner
    from tests.test_v10364 import _replicas
    one,make,docs,original,_=_replicas(monkeypatch)
    payload,images=document_with_images();service=DocumentImport(one.projects)
    progress=service.begin(payload,'Synthetic.docx','heading1','Synthetic images',owner,'op_import_images')
    pid=service._load(owner,progress['job_id'])['project_id']
    assert progress['images']==5 and not progress['warnings']
    assert len(docs.query_collection(('projects',pid,'assets')))==5
    with pytest.raises(ProjectUnavailable):make().projects.get(pid)
    while not progress['ready']:progress=service.step(owner,progress['job_id'],progress['next_batch'])
    result=service.finish(owner,progress['job_id'])
    assert result['images']==5
    rows=make().projects.domain.list('manuscript_block',{'project_id':pid},user_id=owner)
    image_blocks=[row for row in rows if row['block_type']=='image']
    assert len(image_blocks)==5
    with one.context.bind(pid,user_id=owner):
        for block,blob in zip(image_blocks,images):assert one.assets.read_blob(block['image_path'])==blob
    # The same authenticated operation can resume without another publication.
    assert service.begin(payload,'Synthetic.docx','heading1','Synthetic images',owner,'op_import_images')['job_id']==progress['job_id']
    assert service.finish(owner,progress['job_id'])['project']['id']==pid
    assert not docs.query_collection(('projects',original,'assets'))
    with pytest.raises(ValueError,match='ALREADY_PUBLISHED'):service.packages._stage_asset(pid,owner,images[0],'.png')


def test_hosted_image_failure_never_publishes_partial_project(monkeypatch):
    from hosted.document_import import DocumentImport
    from hosted.runtime_services import _DEV_USER_ID as owner
    from tests.test_v10364 import _replicas
    one,_,docs,original,_=_replicas(monkeypatch)
    before={row.path for row in docs.query_collection(('projects',))}
    service=DocumentImport(one.projects);payload,_=document_with_images()
    def failed(*args,**kwargs):raise RuntimeError('Synthetic image publication failed')
    monkeypatch.setattr(service.packages,'_stage_asset',failed)
    with pytest.raises(RuntimeError,match='publication failed'):
        service.begin(payload,'Synthetic.docx','heading1','Synthetic failed images',owner,'op_import_images_failed')
    assert {row.path for row in docs.query_collection(('projects',))}==before
    assert not docs.query_collection(('projects',original,'assets'))
