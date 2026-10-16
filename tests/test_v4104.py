from source_contracts import frontend_source_from_app_path
import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK

from conftest import assert_ok


def make_multitab_docx():
    doc=Document()
    first=doc.add_paragraph('Research',style='Title'); first.add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading('Sources',level=1)
    doc.add_paragraph('Research prose')
    second=doc.add_paragraph('Characters',style='Title'); second.add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading('Cinna',level=1)
    doc.add_paragraph('Character prose')
    out=io.BytesIO(); doc.save(out); return out.getvalue()


def test_v4104_document_import_appends_every_source_tab_after_anchor(client,fresh_project):
    boot=fresh_project['bootstrap']
    manuscript=next(tab for tab in boot['tabs'] if tab['is_manuscript'])
    notes_a=assert_ok(client.post('/api/documents',json={'title':'Notes A'}))
    notes_b=assert_ok(client.post('/api/documents',json={'title':'Notes B'}))

    response=client.post(
        f"/api/documents/import?after_document_id={notes_a['id']}&chapter_heading=heading1",
        files={'file':('multi-tabs.docx',make_multitab_docx(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},
    )
    assert response.status_code==200,response.text
    result=response.json()
    assert result['tabs']==2
    assert len(result['document_ids'])==2

    after=assert_ok(client.get('/api/bootstrap'))
    tabs=after['tabs']
    assert [tab['title'] for tab in tabs]==['Manuscript','Notes A','Research','Characters','Notes B']
    imported=[tab for tab in tabs if tab['id'] in result['document_ids']]
    assert imported and all(not tab['is_manuscript'] for tab in imported)
    assert manuscript['id'] in [tab['id'] for tab in tabs]

    research=next(tab for tab in tabs if tab['title']=='Research')
    characters=next(tab for tab in tabs if tab['title']=='Characters')
    research_chapter=next(ch for ch in after['chapters'] if ch['tab_id']==research['id'])
    characters_chapter=next(ch for ch in after['chapters'] if ch['tab_id']==characters['id'])
    assert research_chapter['title']=='Sources'
    assert characters_chapter['title']=='Cinna'
    research_data=assert_ok(client.get(f"/api/chapters/{research_chapter['id']}"))
    character_data=assert_ok(client.get(f"/api/chapters/{characters_chapter['id']}"))
    assert any(block['text_content']=='Research prose' for block in research_data['blocks'])
    assert any(block['text_content']=='Character prose' for block in character_data['blocks'])


def test_v4104_single_source_tab_uses_filename_instead_of_generic_manuscript(client,fresh_project):
    boot=fresh_project['bootstrap']; anchor=boot['tabs'][0]
    doc=Document(); doc.add_heading('Only Chapter',level=1); doc.add_paragraph('Imported text')
    out=io.BytesIO(); doc.save(out)
    result=assert_ok(client.post(
        f"/api/documents/import?after_document_id={anchor['id']}",
        files={'file':('World Notes.docx',out.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},
    ))
    after=assert_ok(client.get('/api/bootstrap'))
    imported=next(tab for tab in after['tabs'] if tab['id']==result['first_document_id'])
    assert imported['title']=='World Notes'
    assert not imported['is_manuscript']


def test_v4104_static_contract_exposes_import_document_on_tabs_collection_menu():
    root=Path(__file__).resolve().parents[1]
    js=frontend_source_from_app_path(root/'static/app.js')
    html=(root/'static/index.html').read_text(encoding='utf-8')
    assert "label:'Import Tabs…'" in js
    assert 'beginDocumentImport(activeDocument()||state.documents.at(-1))' in js
    assert "label:'Import tabs…'" not in js
    assert '/api/documents/import?after_document_id=' in js
    assert 'Every source tab will be appended after' in js
    assert 'id="documentImportInput"' in html
