from __future__ import annotations

import io
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from PIL import Image
from conftest import assert_ok


def make_docx() -> bytes:
    doc=Document()
    doc.add_heading("Chapter One", level=1)
    p=doc.add_paragraph()
    p.add_run("Tabbed")
    p.add_run("\t")
    hi=p.add_run("wordy")
    hi.font.highlight_color=WD_COLOR_INDEX.YELLOW
    img=io.BytesIO(); Image.new("RGB",(8,8),(255,255,255)).save(img,format="PNG"); img.seek(0)
    doc.add_picture(img)
    doc.add_paragraph("After image")
    out=io.BytesIO(); doc.save(out); return out.getvalue()


def test_docx_import_preserves_tabs_image_order_highlight_and_heading(client, fresh_project):
    response=client.post("/api/import",files={"file":("book.docx",make_docx(),"application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code==200, response.text
    result=response.json(); assert result["images"]>=1 and result["chapters"]>=1
    boot=assert_ok(client.get("/api/bootstrap"))
    ch=next(c for c in boot["chapters"] if c["title"]=="Chapter One")
    data=assert_ok(client.get(f"/api/chapters/{ch['id']}"))
    kinds=[b["block_type"] for b in data["blocks"]]
    assert "image" in kinds
    texts=[b["text_content"] for b in data["blocks"] if b["block_type"]=="paragraph"]
    assert any("Tabbed\twordy" in t for t in texts)
    image_index=kinds.index("image")
    after_index=next(i for i,b in enumerate(data["blocks"]) if b["text_content"]=="After image")
    assert image_index < after_index
    annotations=data["annotations"]
    assert any(a["annotation_type"]=="highlight" and "wordy" in a["selected_text"] and a["label_name"].startswith("Imported") for a in annotations)


def test_clean_docx_export_excludes_ghost_scene_title_and_editor_markup(client, fresh_project):
    ch=fresh_project["bootstrap"]["chapters"][0]
    ids=assert_ok(client.put(f"/api/chapters/{ch['id']}/blocks",json={"blocks":[{"text_content":"Visible prose","content_html":"Visible prose"}]}))["ids"]
    assert_ok(client.post("/api/scenes",json={"title":"SECRET GHOST SCENE","block_ids":ids}))
    marker=assert_ok(client.post("/api/labels",json={"name":"Wordiness","color":"#ffd000"}))
    assert_ok(client.post("/api/annotations/batch",json={"annotation_type":"highlight","label_id":marker["id"],"items":[{"block_id":ids[0],"start_offset":0,"end_offset":7,"selected_text":"Visible"}]}))
    exported=client.get("/api/export/docx")
    assert exported.status_code==200
    doc=Document(io.BytesIO(exported.content))
    text="\n".join(p.text for p in doc.paragraphs)
    assert "Visible prose" in text
    assert "SECRET GHOST SCENE" not in text
