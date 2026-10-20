from __future__ import annotations

import struct
import zlib
from pathlib import Path

import pytest

from security.image_security import (
    HOSTED_DEFAULT_IMAGE_BYTES,
    HOSTED_HARD_MAX_IMAGE_BYTES,
    HOSTED_MAX_IMAGE_DIMENSION,
    HOSTED_MAX_IMAGE_PIXELS,
    HOSTED_SAFE_IMAGE_EXTENSIONS,
    UnsafeImage,
    inspect_hosted_image_bytes,
)

ROOT=Path(__file__).resolve().parents[1]


def _chunk(kind:bytes,payload:bytes)->bytes:
    return struct.pack('>I',len(payload))+kind+payload+struct.pack('>I',zlib.crc32(kind+payload)&0xffffffff)


def _png(width:int,height:int)->bytes:
    ihdr=struct.pack('>IIBBBBB',width,height,8,6,0,0,0)
    return b'\x89PNG\r\n\x1a\n'+_chunk(b'IHDR',ihdr)+_chunk(b'IEND',b'')


def test_v10315_hosted_image_boundary_is_small_and_format_narrow():
    assert tuple(map(int,(ROOT/'VERSION').read_text(encoding='utf-8').strip().split('.'))) >= (10,3,15)
    assert HOSTED_DEFAULT_IMAGE_BYTES==12*1024*1024
    assert HOSTED_HARD_MAX_IMAGE_BYTES==25*1024*1024
    assert HOSTED_MAX_IMAGE_DIMENSION==8192
    assert HOSTED_MAX_IMAGE_PIXELS==20_000_000
    assert HOSTED_SAFE_IMAGE_EXTENSIONS==frozenset({'.png','.jpg','.webp'})


def test_v10315_image_validation_uses_structure_not_filename_or_mime():
    info=inspect_hosted_image_bytes(_png(1024,768),'.png')
    assert (info.extension,info.media_type,info.width,info.height)==('.png','image/png',1024,768)
    with pytest.raises(UnsafeImage,match='IMAGE_FORMAT_MISMATCH'):
        inspect_hosted_image_bytes(_png(100,100),'.jpg')
    with pytest.raises(UnsafeImage,match='IMAGE_FORMAT_UNSUPPORTED'):
        inspect_hosted_image_bytes(b'....ftypavif'+b'0'*100,'.avif')
    with pytest.raises(UnsafeImage,match='IMAGE_BYTES_TOO_LARGE'):
        inspect_hosted_image_bytes(_png(100,100),'.png',byte_limit=32)


def test_v10315_image_geometry_and_png_integrity_fail_closed_without_decoder():
    with pytest.raises(UnsafeImage,match='IMAGE_DIMENSIONS_TOO_LARGE|IMAGE_PIXELS_TOO_LARGE'):
        inspect_hosted_image_bytes(_png(8192,8192),'.png')
    damaged=bytearray(_png(100,100)); damaged[30]^=1
    with pytest.raises(UnsafeImage,match='PNG_CRC_INVALID'):
        inspect_hosted_image_bytes(bytes(damaged),'.png')


def test_v10315_hosted_upload_and_portable_assets_validate_before_publication():
    upload=(ROOT/'src/api/import_export/assets.py').read_text(encoding='utf-8')
    cuts=(ROOT/'src/api/durability/cuts.py').read_text(encoding='utf-8')
    project_assets=(ROOT/'src/api/project_management/assets.py').read_text(encoding='utf-8')
    portable=(ROOT/'src/api/reviews/portable.py').read_text(encoding='utf-8')
    packages=(ROOT/'src/api/project_management/packages.py').read_text(encoding='utf-8')
    assert 'inspect_hosted_image_bytes(payload,ext,byte_limit=limit)' in upload
    assert upload.index('inspect_hosted_image_bytes(payload,ext,byte_limit=limit)') < upload.index('runtime.assets.put_blob(payload,ext)')
    assert 'inspect_hosted_image_bytes(blob,ext,byte_limit=_HOSTED_LIMITS.image_bytes)' in cuts
    assert 'zi.file_size>image_limit' in cuts
    assert 'inspect_hosted_image_bytes(payload,suffix,byte_limit=_HOSTED_LIMITS.image_bytes)' in project_assets
    assert project_assets.index('inspect_hosted_image_bytes(payload,suffix,byte_limit=_HOSTED_LIMITS.image_bytes)') < project_assets.index('return Response(content=payload')
    assert 'inspect_hosted_image_bytes(data,Path(name).suffix,byte_limit=_HOSTED_LIMITS.image_bytes)' in portable
    assert 'max_asset_bytes=(_HOSTED_LIMITS.image_bytes if _HOSTED_LIMITS is not None else breadreview.MAX_ASSET_BYTES)' in portable
    assert 'HOSTED_PROJECT_IMAGE_UNSAFE' in packages and 'inspect_hosted_image_bytes(archive.read(info)' in packages
    assert "'X-Content-Type-Options':'nosniff'" in project_assets


def test_v10315_hosted_server_never_decodes_project_images_in_process():
    pdf=(ROOT/'src/api/import_export/pdf_export.py').read_text(encoding='utf-8')
    docx=(ROOT/'src/api/import_export/docx_export.py').read_text(encoding='utf-8')
    assert '_HOSTED_LIMITS is not None' in pdf and "RLParagraph('[Image]'" in pdf
    assert '_HOSTED_LIMITS is not None' in docx and "add_run('[Image]')" in docx
    # Local export retains image support through the environment-neutral bytes port;
    # hosted export reaches the placeholder branch before any in-process decoder.
    assert 'RLImage(BytesIO(runtime.assets.read_blob(asset_name)))' in pdf
    assert '.add_picture(BytesIO(runtime.assets.read_blob(asset_name))' in docx
    assert pdf.index("if _HOSTED_LIMITS is not None") < pdf.index('RLImage(BytesIO(runtime.assets.read_blob(asset_name)))')
    assert docx.index("if _HOSTED_LIMITS is not None") < docx.index('.add_picture(BytesIO(runtime.assets.read_blob(asset_name))')


def test_v10315_container_excludes_broad_heif_imagemagick_stack_and_pins_pillow():
    requirements=(ROOT/'requirements.txt').read_text(encoding='utf-8')
    docker=(ROOT/'deploy/Dockerfile.hosted-dev').read_text(encoding='utf-8')
    guard=(ROOT/'scripts/assert_hosted_image_stack.py').read_text(encoding='utf-8')
    assert 'Pillow==13.0.0' in requirements
    assert 'FROM python:3.14.7-slim-trixie' in docker
    assert 'RUN python scripts/assert_hosted_image_stack.py' in docker
    for token in ("('magick','convert')","find_library('heif')","find_spec('pillow_heif')"):
        assert token in guard


def test_patched_pillow_keeps_pdf_and_docx_image_export_compatible():
    from io import BytesIO
    from PIL import Image
    from docx import Document
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen.canvas import Canvas

    for kind in ('PNG', 'JPEG', 'WEBP'):
        payload = BytesIO()
        Image.new('RGB', (8, 6), '#226644').save(payload, format=kind)
        raw = payload.getvalue()
        reader = ImageReader(BytesIO(raw))
        assert reader.getSize() == (8, 6)
        pdf = BytesIO()
        canvas = Canvas(pdf)
        canvas.drawImage(reader, 0, 0, width=8, height=6)
        canvas.save()
        assert pdf.getvalue().startswith(b'%PDF-')
        if kind in ('PNG', 'JPEG'):
            document = Document()
            document.add_picture(BytesIO(raw))
            package = BytesIO()
            document.save(package)
            assert package.getvalue().startswith(b'PK')


def test_v10315_frontend_narrows_hosted_images_without_regressing_local_picker():
    index=(ROOT/'static/index.html').read_text(encoding='utf-8')
    boot=(ROOT/'static/js/core/boot.js').read_text(encoding='utf-8')
    annotations=(ROOT/'static/js/features/manuscript/annotations.js').read_text(encoding='utf-8')
    accept=index.split('id="manuscriptImageFileInput"',1)[1].split('>',1)[0]
    # Local mode preserves its historical browser-supported picker surface.
    for allowed in ('image/png','image/jpeg','image/webp','image/gif','image/avif','image/bmp'):
        assert allowed in accept.lower()
    # Hosted mode narrows both file pickers and drag/drop after runtime attestation.
    assert "const HOSTED_SAFE_IMAGE_ACCEPT='image/png,image/jpeg,image/webp,.png,.jpg,.jpeg,.webp'" in boot
    assert "runtimeSupports('hosted_projects')" in boot
    assert "if(canvas)canvas.accept=hosted?HOSTED_SAFE_IMAGE_ACCEPT:'image/*'" in boot
    assert "['image/png','image/jpeg','image/webp'].includes(type)" in annotations
    assert '(png|jpe?g|webp|gif|avif|bmp)' in annotations

