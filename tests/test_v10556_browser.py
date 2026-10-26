"""Full Text previews are responsive but never publish partial editable state."""
import io
from docx import Document
from playwright.sync_api import expect
from tests.test_v10548_browser import hosted_browser


def _project(base,page):
    doc=Document()
    for n in range(2):
        doc.add_heading(f'Synthetic56 chapter {n}',level=1)
        for i in range(260):doc.add_paragraph(f'Synthetic56 chapter {n} paragraph {i}')
    raw=io.BytesIO();doc.save(raw);headers={'Origin':base}
    progress=page.request.post(base+'/api/import/new-project?resumable=true&project_name=Synthetic56',
        multipart={'file':{'name':'Synthetic56.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','buffer':raw.getvalue()}},headers=headers).json()
    job=progress['job_id']
    while not progress['ready']:progress=page.request.post(base+f'/api/import/jobs/{job}/step?batch={progress["next_batch"]}',headers=headers).json()
    pid=page.request.post(base+f'/api/import/jobs/{job}/finish',headers=headers).json()['project']['id']
    page.goto(base+'/p/'+pid,wait_until='domcontentloaded');page.wait_for_function('()=>document.documentElement.dataset.breadmakerBootReady==="true"')
    return pid


def test_full_text_preview_is_read_only_until_final_validation_then_all_text_is_editable(tmp_path):
    with hosted_browser(tmp_path) as (base,page):
        _project(base,page);held=[]
        def hold_final(route):
            if route.request.headers.get('if-none-match'):held.append(route)
            else:route.continue_()
        page.route('**/api/book?**',hold_final)
        page.locator('#bookViewBtn').click();expect(page.locator('.full-text-preview')).to_be_visible()
        assert page.locator('.full-text-preview').evaluate('(el)=>el.matches("#documentScroller > article.paper")')
        assert page.locator('.full-text-preview button,.full-text-preview header').count()==0
        expect(page.locator('.manuscript-loading-indicator')).to_be_visible()
        assert 'Loading Full Text' not in page.locator('#saveState').inner_text()
        expect(page.locator('.full-text-preview-body')).to_contain_text('Synthetic56 chapter 1 paragraph 259')
        assert page.evaluate('()=>state.viewMode')=='chapter'
        assert page.locator('.full-text-preview [contenteditable="true"]').count()==0
        assert page.locator('.manuscript-block').count()==260
        page.wait_for_function('()=>document.querySelector(".full-text-preview-body")?.textContent.includes("paragraph 259")')
        assert held;held[0].continue_()
        page.wait_for_function('()=>state.viewMode==="book"')
        expect(page.locator('.full-text-preview')).to_have_count(0)
        expect(page.locator('.manuscript-loading-indicator')).to_have_count(0)
        assert page.locator('.manuscript-block').count()==520
        assert not page.evaluate('()=>state.dirty')


def test_cancelled_or_changed_full_text_never_replaces_current_chapter(tmp_path):
    with hosted_browser(tmp_path) as (base,page):
        _project(base,page);old=page.locator('#editor').inner_text();held=[]
        def hold_final(route):
            if route.request.headers.get('if-none-match'):held.append(route)
            else:route.continue_()
        page.route('**/api/book?**',hold_final)
        page.locator('#bookViewBtn').click()
        expect(page.locator('.full-text-preview-body')).to_contain_text('Synthetic56 chapter 1 paragraph 259')
        page.locator('.full-text-preview').press('Escape')
        expect(page.locator('.full-text-preview')).to_have_count(0)
        assert held;held[0].continue_()
        page.wait_for_function('()=>!fullTextLoading')
        assert page.evaluate('()=>state.viewMode')=='chapter' and page.locator('#editor').inner_text()==old
        assert not page.evaluate('()=>state.dirty')
        # A changed final revision must fail just as safely, without publishing
        # the preview or caching a partial book as the editable manuscript.
        held.clear();page.locator('#bookViewBtn').click()
        expect(page.locator('.full-text-preview-body')).to_contain_text('Synthetic56 chapter 1 paragraph 259')
        assert held;held[0].fulfill(status=200,json={'project':{'generation':999999}})
        page.wait_for_function('()=>!fullTextLoading')
        expect(page.locator('.full-text-preview')).to_have_count(0)
        assert page.evaluate('()=>state.viewMode')=='chapter' and page.locator('#editor').inner_text()==old
        assert page.evaluate('()=>state.bookCache.size')==0
